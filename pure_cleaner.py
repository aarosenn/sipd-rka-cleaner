import os
import re
import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def bersihkan_rka_pure_python(source_path, output_path):
    """
    Versi 100% Python Murni (tanpa PowerShell/Excel COM).
    Memerlukan pustaka openpyxl dan xlrd.
    """
    if not os.path.exists(source_path):
        print(f"[ERROR] File tidak ditemukan: {source_path}")
        return

    wb_src = xlrd.open_workbook(source_path)
    sh_src = wb_src.sheet_by_index(0)

    items = []
    for r in range(44, sh_src.nrows):
        rek_val = str(sh_src.cell_value(r, 1)).strip()
        uraian_val = str(sh_src.cell_value(r, 3)).strip()
        sblm_val = sh_src.cell_value(r, 14)
        ssdh_val = sh_src.cell_value(r, 23)

        if re.match(r'^[0-9]+(\.[0-9]+)*$', rek_val):
            try:
                val_sblm = float(str(sblm_val).replace('.', '').replace(',', '.'))
            except (ValueError, TypeError):
                val_sblm = 0.0
                
            try:
                val_ssdh = float(str(ssdh_val).replace('.', '').replace(',', '.'))
            except (ValueError, TypeError):
                val_ssdh = 0.0

            items.append({
                "kode": rek_val,
                "uraian": uraian_val,
                "sebelum": val_sblm,
                "sesudah": val_ssdh
            })

    print(f"[OK] Berhasil mengekstrak {len(items)} rekening belanja.")

    wb_new = Workbook()
    ws = wb_new.active
    ws.title = "RKA_Kompak"
    ws.views.sheetView[0].showGridLines = True

    # Font & Styling
    font_title = Font(name="Segoe UI", size=13, bold=True)
    font_sub = Font(name="Segoe UI", size=10.5, bold=True)
    font_hdr = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=9.5, bold=True)
    font_regular = Font(name="Segoe UI", size=9.5)

    fill_hdr = PatternFill(start_color="1F365D", end_color="1F365D", fill_type="solid")
    fill_lvl1 = PatternFill(start_color="F0D5C7", end_color="F0D5C7", fill_type="solid")
    fill_lvl2 = PatternFill(start_color="FAE8DF", end_color="FAE8DF", fill_type="solid")
    fill_lvl3 = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='EDEDED'),
        bottom=Side(style='thin', color='EDEDED')
    )

    ws["A1"] = "PEMERINTAH KABUPATEN BLITAR"
    ws["A1"].font = font_title
    ws["A2"] = "RENCANA KERJA DAN ANGGARAN (RKA-BELANJA PERUBAHAN) - TAHUN ANGGARAN 2026"
    ws["A2"].font = font_sub
    ws["A4"] = "SKPD: Dinas Komunikasi, Informatika, Statistik dan Persandian (2.16.2.20.2.21.01.0000)"
    ws["A5"] = "Sub Kegiatan: 2.16.01.2.02.0001 Penyediaan Gaji dan Tunjangan ASN"
    ws["A6"] = "Sumber Dana: Dana Alokasi Umum (DAU) | Target: 37 Orang / bulan"

    headers = [
        "Kode Rekening", "Uraian Belanja", "Sumber Dana",
        "Sebelum Perubahan (Rp)", "Sesudah Perubahan (Rp)",
        "Bertambah / (Berkurang) (Rp)", "% Perubahan"
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=8, column=col_idx, value=h)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[8].height = 26

    curr_row = 9
    for it in items:
        dots = it["kode"].count('.')
        
        c1 = ws.cell(row=curr_row, column=1, value=it["kode"])
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c1.number_format = "@"

        indent_spaces = "    " * min(dots, 5)
        c2 = ws.cell(row=curr_row, column=2, value=indent_spaces + it["uraian"])
        c2.alignment = Alignment(horizontal="left", vertical="center")

        c3 = ws.cell(row=curr_row, column=3, value="DAU" if dots >= 4 else "")
        c3.alignment = Alignment(horizontal="center", vertical="center")

        c4 = ws.cell(row=curr_row, column=4, value=it["sebelum"])
        c4.number_format = '#,##0'
        c4.alignment = Alignment(horizontal="right", vertical="center")

        c5 = ws.cell(row=curr_row, column=5, value=it["sesudah"])
        c5.number_format = '#,##0'
        c5.alignment = Alignment(horizontal="right", vertical="center")

        c6 = ws.cell(row=curr_row, column=6, value=f"=E{curr_row}-D{curr_row}")
        c6.number_format = '#,##0;[Red]-#,##0;"-"'
        c6.alignment = Alignment(horizontal="right", vertical="center")

        c7 = ws.cell(row=curr_row, column=7, value=f"=IF(D{curr_row}=0,0,(E{curr_row}-D{curr_row})/D{curr_row})")
        c7.number_format = '0.00%;[Red]-0.00%;0.00%'
        c7.alignment = Alignment(horizontal="right", vertical="center")

        for col_idx in range(1, 8):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.border = thin_border
            if dots == 0:
                cell.font = font_bold
                cell.fill = fill_lvl1
            elif dots == 1:
                cell.font = font_bold
                cell.fill = fill_lvl2
            elif dots == 2:
                cell.font = font_bold
                cell.fill = fill_lvl3
            elif dots in (3, 4):
                if col_idx <= 2: cell.font = font_bold
            else:
                cell.font = font_regular
                
        ws.row_dimensions[curr_row].height = 20
        curr_row += 1

    # Total Row
    ws.merge_cells(f"A{curr_row}:C{curr_row}")
    tot_label = ws.cell(row=curr_row, column=1, value="TOTAL BELANJA DAERAH (KODE REKENING 5)")
    tot_label.alignment = Alignment(horizontal="right", vertical="center")
    tot_label.font = font_bold

    t4 = ws.cell(row=curr_row, column=4, value="=D9")
    t4.number_format = '#,##0'
    t5 = ws.cell(row=curr_row, column=5, value="=E9")
    t5.number_format = '#,##0'
    t6 = ws.cell(row=curr_row, column=6, value="=F9")
    t6.number_format = '#,##0;[Red]-#,##0;"-"'
    t7 = ws.cell(row=curr_row, column=7, value="=G9")
    t7.number_format = '0.00%;[Red]-0.00%;0.00%'

    for col in range(1, 8):
        c = ws.cell(row=curr_row, column=col)
        c.font = font_bold
        c.fill = PatternFill(start_color="E8D0C0", end_color="E8D0C0", fill_type="solid")

    # Lebar Kolom Acuan Digit Terpanjang
    col_widths = {1: 24.0, 2: 50.0, 3: 14.0, 4: 22.0, 5: 22.0, 6: 24.0, 7: 14.0}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A9"
    wb_new.save(output_path)
    print(f"[SELESAI] File berhasil disimpan ke:\n{output_path}")

if __name__ == "__main__":
    src = "contoh_rka.xls"
    dst = "RKA_COMPACT.xlsx"
    bersihkan_rka_pure_python(src, dst)
